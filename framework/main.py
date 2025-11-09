import sys, os
from docx import Document 
import numpy as np
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from framework.signals import (
    generate_signal,
    Signal
)
from framework.fileHandling import load_signal,save_signal
from framework.operations import (
    subtract_signals,
    normalize_signal,
    quantization,
    fourier_transform,
    remove_dc_component
)
from tests.signalcompare import (
    SignalComapreAmplitude,
    SignalComaprePhaseShift
)
from tests.CompareSignals import (
    SignalsAreEqual
)
def read_dft_file(path):
    """
    Reads a DFT output text file with this format:
        1
        0
        8
        idx || amp || phase
        0 64 0
        1 20.9050074380220 1.96349540849362
        ...

    Returns:
        tuple(list[float], list[float]):
            (amplitudes, phases)
    """
    amplitudes, phases = [], []
    with open(path, 'r') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        if not line or line.startswith("idx") or "||" in line:
            continue

        parts = line.replace("f", "").replace(",", "").split()
        if len(parts) < 3:
            try:
                _ = int(parts[0])
                continue
            except Exception:
                continue

        try:
            _, amp_str, ph_str = parts[:3]
            amplitudes.append(float(amp_str))
            phases.append(float(ph_str))
        except ValueError:
            continue

    return amplitudes, phases

def read_txt(path):
    """
    Reads amplitude and phase values from a plain text file.
    Automatically cleans C-style 'f' suffixes or stray commas.
    Returns two lists of floats.
    """
    amplitudes, phases = [], []
    with open(path, 'r') as f:
        for line in f:
            parts = line.strip().split()
            if len(parts) >= 2:
                amp_str = parts[0].strip().replace(',', '').replace('f', '')
                ph_str = parts[1].strip().replace(',', '').replace('f', '')
                try:
                    amplitudes.append(float(amp_str))
                    phases.append(float(ph_str))
                except ValueError:
                    continue
    return amplitudes, phases



def read_output_file(path):
    """
    Reads a DOCX file containing amplitude and phase values.
    Expects each row to have two numeric columns: amplitude and phase.
    Returns two lists: amplitudes, phases.
    """
    doc = Document(path)
    amplitudes, phases = [], []

    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) >= 2:
                try:
                    amp = float(cells[0].text.strip())
                    ph = float(cells[1].text.strip())
                    amplitudes.append(amp)
                    phases.append(ph)
                except ValueError:
                    # skip rows that don't contain numeric data
                    continue

    if not amplitudes and not phases:
        for para in doc.paragraphs:
            parts = para.text.strip().split()
            if len(parts) >= 2:
                try:
                    amplitudes.append(float(parts[0]))
                    phases.append(float(parts[1]))
                except ValueError:
                    continue

    return amplitudes, phases



DFT_input = load_signal("/home/mohammed/repos/dsp/Inputs/input_Signal_DFT.txt")
dft_result = fourier_transform(DFT_input,False)  # Should produce frequency-domain representation
t_amp,t_phase = read_txt("/home/mohammed/repos/dsp/outputs/output_signal_DFT.txt")
amplitude_test = SignalComapreAmplitude(dft_result.y,t_amp)
phase_test = SignalComaprePhaseShift(dft_result.phase,t_phase)

########################################################
# IDFT

ny,nf = read_dft_file("/home/mohammed/repos/dsp/Inputs/input_signal_IDFT.txt")

idft_input = Signal( 
    name="Signal",
    signal_type=1,
    x=list(range(len(ny))),
    y=ny,
    phase=nf
)

trans = fourier_transform(idft_input,True)

out = load_signal("/home/mohammed/repos/dsp/outputs/Output_Signal_IDFT.txt")

a2 = SignalComapreAmplitude(trans.y,out.y)
# print(trans.y)
# print(out.y)
if a2 and amplitude_test and phase_test:
    print("Tests Passed Successfully")
else:
    if not amplitude_test:
        print("Amplitude comparison failed")
    if not phase_test:
        print(" Phase comparison failed")
    if not a2 :
        print("IDFT Amplitude Failed")

